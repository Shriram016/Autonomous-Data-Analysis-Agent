"""
check_compare.py — Verifies the comparison workbook built from the two REAL runs. Offline, free, no LLM calls.

Builds the table into a temp file and checks that it is internally consistent and agrees with the run folders:
63 rows, every query once, match/error flags agree with the outcome column, every failure has a reason,
and the strict scores equal the pass counts recorded in each run's manifest.

Run from the project root:
    python llm_codegen_experiment/checks/check_compare.py
"""

import json
import os
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from llm_codegen_experiment import compare as cmp  # noqa: E402

ok = True


def check(name, cond, detail=""):
    global ok
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f" — {detail}" if detail and not cond else ""), flush=True)
    ok &= bool(cond)


def main() -> int:
    llm_dir = cmp._latest(os.path.join(cmp.LLM_RESULTS, "full_*"))
    adaa_dir = cmp.DEFAULT_ADAA_RUN
    out = os.path.join(tempfile.mkdtemp(), "cmp.xlsx")
    code = cmp.main(["--adaa", adaa_dir, "--llm", llm_dir, "--out", out])
    check("comparison builds", code == 0 and os.path.exists(out))

    wb = load_workbook(out)
    ws = wb["comparison"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    reg = cmp.case_registry()

    check("3 sheets: comparison, summary, notes", wb.sheetnames == ["comparison", "summary", "notes"])
    check("63 rows, each query exactly once", len(rows) == 63 and {r["query_id"] for r in rows} == set(reg))
    wanted = ["query_id", "query", "category", "ground_truth_value", "adaa_value", "llm_value", "gt_vs_adaa", "gt_vs_llm",
              "adaa_error", "llm_error", "adaa_error_reason", "llm_error_reason"]
    check("all requested columns present", all(c in header for c in wanted), str([c for c in wanted if c not in header]))
    check("match and error columns are real booleans",
          all(isinstance(r[c], bool) for r in rows for c in ("gt_vs_adaa", "gt_vs_llm", "adaa_error", "llm_error")))

    for s in ("adaa", "llm"):
        good = {"correct", "refused_correctly"}
        check(f"{s}: match flag agrees with the outcome", all(r[f"gt_vs_{s}"] == (r[f"{s}_outcome"] in good) for r in rows))
        check(f"{s}: error flag is true only for outcome 'error'", all(r[f"{s}_error"] == (r[f"{s}_outcome"] == "error") for r in rows))
        check(f"{s}: error and match are never both true", not any(r[f"{s}_error"] and r[f"gt_vs_{s}"] for r in rows))
        check(f"{s}: every non-match has a reason", all(r[f"{s}_error_reason"] for r in rows if not r[f"gt_vs_{s}"]))
        check(f"{s}: a wrong answer is not counted as an error",
              any(r[f"{s}_outcome"] == "wrong_answer" for r in rows) and all(not r[f"{s}_error"] for r in rows if r[f"{s}_outcome"] == "wrong_answer"))

    passed_adaa = json.load(open(os.path.join(adaa_dir, "manifest.json"), encoding="utf-8"))["totals"]["passed"]
    passed_llm = json.load(open(os.path.join(llm_dir, "manifest.json"), encoding="utf-8"))["totals"]["passed"]
    check(f"ADAA strict score equals its run's pass count ({passed_adaa})", sum(r["adaa_strict_match"] for r in rows) == passed_adaa)
    check(f"code system strict score equals its run's pass count ({passed_llm})", sum(r["llm_strict_match"] for r in rows) == passed_llm)
    check("name-insensitive score is never lower than strict",
          sum(r["gt_vs_adaa"] for r in rows) >= passed_adaa and sum(r["gt_vs_llm"] for r in rows) >= passed_llm)

    refuse = [r for r in rows if r["expected_behavior"] == "refuse"]
    check("10 compound queries expect a refusal", len(refuse) == 10 and all(r["ground_truth_value"].startswith("(correct = refuse") for r in refuse))
    print("\nALL CHECKS PASSED" if ok else "\nSOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
