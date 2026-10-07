"""Offline tests for the comparison table builder (llm_codegen_experiment/compare.py), on synthetic records."""

import pytest
from openpyxl import load_workbook

from llm_codegen_experiment import compare as cmp

META = {"query": "q", "category": "c", "compare_mode": "ordered", "float_tol": 0.01, "kind": "single"}
GT = [{"State": "California", "Sales_sum": 457687.63}, {"State": "New York", "Sales_sum": 310876.27}]


def rec(status="success", data=None, gt=GT, expected="answer", message="", sandbox="ok", **kw):
    r = {"id": "Q11", "kind": "single", "expected_behavior": expected, "pipeline_status": status,
         "pipeline_message": message, "pipeline_data": data, "gt_data": gt, "value_match": False, "passed": False,
         "trace": [{"sandbox_status": sandbox}], "mismatches": [], "cost_usd": 0.0001, "duration_s": 1.0}
    r.update(kw)
    return r


# ---------------------------------------------------------------- value comparison ignores column names, not values or order
def test_column_names_and_column_order_are_ignored():
    renamed = [{"Sales": 457687.63, "State": "California"}, {"Sales": 310876.27, "State": "New York"}]
    assert cmp.name_insensitive_match(rec(data=renamed), META) is True


def test_values_and_row_order_still_matter():
    swapped = [GT[1], GT[0]]
    assert cmp.name_insensitive_match(rec(data=swapped), META) is False
    other = [{"s": "California", "v": 457687.63}, {"s": "New York", "v": 999.0}]
    assert "row 2" in cmp.first_difference(GT, other, 0.01)
    assert cmp.first_difference(GT, GT[:1], 0.01) == "1 rows returned, ground truth has 2"


def test_tolerance_is_one_percent_relative():
    close = [{"a": "california", "b": 457687.63 * 1.005}, {"a": "new york", "b": 310876.27}]
    far = [{"a": "california", "b": 457687.63 * 1.05}, {"a": "new york", "b": 310876.27}]
    assert cmp.first_difference(GT, close, 0.01) is None
    assert cmp.first_difference(GT, far, 0.01) is not None


def test_a_case_the_original_comparator_accepted_stays_correct_and_one_it_rejected_is_rechecked_on_rows():
    meta = {**META, "compare_mode": "value_only"}
    assert cmp.name_insensitive_match(rec(data=[], value_match=True), meta) is True          # accepted before: stays correct
    assert cmp.name_insensitive_match(rec(data=GT, value_match=False), meta) is True         # same rows: correct
    different = [{"a": "x", "b": 1.0}, {"a": "y", "b": 2.0}]
    assert cmp.name_insensitive_match(rec(data=different, value_match=False), meta) is False


def test_render_value_shapes():
    assert cmp.render_value([{"x": 2297200.8603}]) == "2,297,200.86"
    assert cmp.render_value(None) == "" and cmp.render_value([]) == ""
    text = cmp.render_value([{"a": i} for i in range(8)])
    assert text.count("\n") == 5 and text.endswith("(+3 more rows)")
    assert "State=California | Sales_sum=457,687.63" in cmp.render_value(GT)


# ---------------------------------------------------------------- the six outcomes; error and wrong answer are different
def test_correct_wrong_answer_and_error_are_three_different_things():
    ok = cmp.classify(rec(data=GT), META, "llm")
    wrong = cmp.classify(rec(data=[{"a": "x", "b": 1.0}, {"a": "y", "b": 2.0}]), META, "llm")
    err = cmp.classify(rec(status="error", data=None, message="blocked: not allowed: name 'open'", sandbox="blocked"), META, "llm")
    assert (ok["outcome"], ok["match"], ok["is_error"]) == ("correct", True, False)
    assert (wrong["outcome"], wrong["match"], wrong["is_error"]) == ("wrong_answer", False, False)   # NOT an error
    assert (err["outcome"], err["match"], err["is_error"]) == ("error", False, True)
    assert "[blocked]" in err["reason"] and "wrong value" in wrong["reason"]


def test_refusals():
    right = cmp.classify(rec(status="unsolvable", data=None, gt=None, expected="refuse", message="two results"), META, "llm")
    wrong = cmp.classify(rec(status="unsolvable", data=None, message="cannot"), META, "llm")
    answered = cmp.classify(rec(data=[{"v": 1}], gt=None, expected="refuse"), META, "llm")
    assert (right["outcome"], right["match"], right["is_error"]) == ("refused_correctly", True, False)
    assert (wrong["outcome"], wrong["match"], wrong["is_error"]) == ("refused_wrongly", False, False)
    assert (answered["outcome"], answered["match"], answered["is_error"]) == ("answered_should_refuse", False, False)
    assert right["value"].startswith("(refused:")


def test_a_skipped_multiturn_case_is_an_error_with_its_reason():
    skipped = {"id": "MT13", "kind": "multi", "expected_behavior": "answer", "outcome": "skipped", "pipeline_status": None,
               "mismatches": ["Skipped: turn 3 returned status=error"], "passed": False}
    out = cmp.classify(skipped, META, "llm")
    assert out["is_error"] is True and out["reason"].startswith("context turn failed")


def test_adaa_reason_carries_the_part_c_category():
    wrong = cmp.classify(rec(data=[{"a": "x", "b": 1.0}, {"a": "y", "b": 2.0}]), META, "adaa", "Planner: missing filter")
    assert wrong["reason"].startswith("Planner: missing filter | wrong value")


# ---------------------------------------------------------------- categories and the full table
def test_registry_covers_all_63_queries_with_sensible_categories():
    reg = cmp.case_registry()
    assert len(reg) == 63 and sum(m["kind"] == "multi" for m in reg.values()) == 16
    assert reg["Q01"]["category"] == "1. Simple aggregation" and reg["Q30"]["category"] == "6. Derived calculations"
    assert reg["Q31"]["category"] == "7. Pseudo-compound: breakdown" and reg["Q36"]["category"] == "7. Pseudo-compound: comparison"
    assert reg["Q38"]["category"] == "8. Compound: two scalars"
    assert reg["MT03"]["category"] == "9. Multi-turn (2 turns)" and reg["MT13"]["category"] == "9. Multi-turn (4 turns)"


def _synthetic_runs(system_right):
    """A record per query: single-turn answerable ones correct, compound ones refused, multi-turn correct."""
    reg = cmp.case_registry()
    recs = []
    for cid, meta in reg.items():
        refuse = cid in {f"Q{i}" for i in range(38, 48)}
        base = {"id": cid, "kind": meta["kind"], "expected_behavior": "refuse" if refuse else "answer",
                "pipeline_status": "unsolvable" if refuse else "success", "pipeline_message": "needs two" if refuse else "",
                "gt_data": None if refuse else [{"v": 1.0}], "pipeline_data": None if refuse else [{"x": 1.0}],
                "value_match": system_right, "passed": system_right, "trace": [{"sandbox_status": "ok"}],
                "mismatches": [], "cost_usd": 0.001, "duration_s": 2.0, "input_tokens": 100, "output_tokens": 10, "llm_call_count": 1}
        recs.append(base)
    return recs


def test_build_rows_and_excel_on_a_full_synthetic_run(tmp_path):
    adaa, llm = _synthetic_runs(True), _synthetic_runs(True)
    rows = cmp.build_rows(adaa, llm)
    assert len(rows) == 63 and [r["query_id"] for r in rows][:2] == ["Q01", "Q02"] and rows[-1]["query_id"] == "MT16"
    assert all(r["gt_vs_adaa"] and r["gt_vs_llm"] for r in rows)
    summary = cmp.build_summary(rows, adaa, llm)
    assert summary["n"] == 63 and summary["n_answer"] == 53 and summary["n_refuse"] == 10
    assert summary["agreement"]["both correct"] == 63

    out = tmp_path / "cmp.xlsx"
    cmp.write_excel(str(out), rows, summary, ["NOTES", "line"])
    wb = load_workbook(out)
    assert wb.sheetnames == ["comparison", "summary", "notes"]
    ws = wb["comparison"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["query_id", "query"] and "gt_vs_adaa" in header and "llm_error_reason" in header
    assert ws.max_row == 64 and ws.freeze_panes == "C2" and ws.auto_filter.ref.startswith("A1:")
    first = {h: c.value for h, c in zip(header, ws[2])}
    assert first["gt_vs_adaa"] is True and first["adaa_error"] is False            # real booleans, not text


def test_summary_counts_a_wrong_system_correctly():
    adaa, llm = _synthetic_runs(True), _synthetic_runs(False)
    for r in llm:                                       # make the "wrong" system really wrong on answerable queries
        if r["expected_behavior"] == "answer":
            r["pipeline_data"] = [{"x": 999.0}]
    rows = cmp.build_rows(adaa, llm)
    s = cmp.build_summary(rows, adaa, llm)
    assert s["match_answerable"] == {"adaa": 53, "llm": 0}
    assert s["answerable"]["llm"]["wrong_answer"] == 53 and s["answerable"]["llm"]["error"] == 0
    assert s["agreement"]["only ADAA correct"] == 53 and s["agreement"]["both correct"] == 10


# ---------------------------------------------------------------- extra rows are allowed (user decision), missing rows are not
REGIONS = [{"Region": "Central", "Sales_sum": 501239.89}, {"Region": "East", "Sales_sum": 678781.24},
           {"Region": "South", "Sales_sum": 391721.91}, {"Region": "West", "Sales_sum": 725457.82}]
VO = {**META, "compare_mode": "value_only"}


def test_four_correct_rows_plus_a_total_row_is_correct():
    answer = [{"R": r["Region"], "Sales": r["Sales_sum"]} for r in REGIONS] + [{"R": "Total", "Sales": 2297200.86}]
    r = rec(data=answer, gt=REGIONS, value_match=False)                      # the strict comparator said no (5 rows vs 4)
    assert cmp.name_insensitive_match(r, VO) is True
    assert cmp.name_insensitive_match(r, META) is True                       # also in an ordered case: GT rows in order, extra row last
    assert cmp.extra_rows(r) == 1


def test_a_missing_ground_truth_row_is_still_wrong():
    answer = [{"R": r["Region"], "Sales": r["Sales_sum"]} for r in REGIONS[:3]] + [{"R": "Total", "Sales": 2297200.86}]
    assert cmp.name_insensitive_match(rec(data=answer, gt=REGIONS, value_match=False), VO) is False


def test_a_wrong_value_among_extra_rows_is_still_wrong():
    answer = [{"R": "Central", "Sales": 1.0}] + [{"R": r["Region"], "Sales": r["Sales_sum"]} for r in REGIONS[1:]] + [{"R": "Total", "Sales": 5.0}]
    assert cmp.name_insensitive_match(rec(data=answer, gt=REGIONS, value_match=False), VO) is False


def test_row_order_is_kept_for_ranked_questions_but_free_for_value_only():
    shuffled = [{"a": r["Region"], "b": r["Sales_sum"]} for r in reversed(REGIONS)]
    r = rec(data=shuffled, gt=REGIONS, value_match=False)
    assert cmp.name_insensitive_match(r, META) is False                      # ordered: reversed ranking is wrong
    assert cmp.name_insensitive_match(r, VO) is True                         # value_only: order does not matter


def test_the_extra_rows_column_and_summary_count_only_correct_answers_with_extras():
    adaa, llm = _synthetic_runs(True), _synthetic_runs(True)
    for r in llm:                                                            # give the code system one extra row on every answerable query
        if r["expected_behavior"] == "answer":
            r["gt_data"] = [{"v": 1.0}]
            r["pipeline_data"] = [{"x": 1.0}, {"x": 2.0}]
            r["value_match"] = False
    rows = cmp.build_rows(adaa, llm)
    assert all(r["gt_vs_llm"] for r in rows) and sum(r["llm_extra_rows"] > 0 for r in rows) == 53
    assert all(r["adaa_extra_rows"] == 0 for r in rows)
    assert cmp.build_summary(rows, adaa, llm)["correct_thanks_to_extra_rows"] == {"adaa": 0, "llm": 53}
