"""
compare.py -- ADAA vs the code-writing system, one row per query, saved as an Excel file.

Reads two finished full-eval run folders (records.jsonl) and builds llm_codegen_experiment/results/
comparison_adaa_vs_llm.xlsx with three sheets:
    comparison  63 rows: category, ground truth, both systems' values, match flags, error flags, reasons
    summary     headline counts, by category, who beats whom, cost and time
    notes       definitions, source runs and caveats

Scoring rule used for the match columns (applied to BOTH systems): every ground-truth row must be present
in the answer, compared with the same 1% relative tolerance as the eval comparator; EXTRA ROWS ARE ALLOWED
(e.g. 4 correct region rows plus a Total row is correct), and column NAMES AND COLUMN ORDER ARE IGNORED (the
code system names columns naturally, e.g. `Sales`, ADAA's tools produce `Sales_sum`). Row order must be kept
for ranked ("ordered"/"full") questions, and is free for "value_only" ones. A case the original comparator
already accepted stays correct. The strict harness score is kept in extra columns.

error vs wrong answer are different things:
    error         the system produced no usable answer (crash, blocked by the sandbox, tool error, skipped)
    wrong answer  it produced a table, but the values do not match ground truth
    A refusal is neither: correct when the question has no ground truth (it needs two separate results),
    a "wrong refusal" when the question was answerable.

Usage (from the project root; offline, no LLM calls, no cost):
    python llm_codegen_experiment/compare.py
    python llm_codegen_experiment/compare.py --adaa eval/results/full_<ts> --llm llm_codegen_experiment/results/full_<ts>
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import sys
from collections import Counter
from typing import Any, Dict, List, Optional

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from eval import failure_labels as fl                                   # noqa: E402
from eval.multiturn_test_cases import MULTITURN_TEST_CASES               # noqa: E402
from eval.run_full_eval import final_records, load_jsonl                 # noqa: E402
from eval.test_cases import TEST_CASES                                   # noqa: E402
from eval.test_cases2 import TEST_CASES_2                                # noqa: E402
from eval.test_cases3 import TEST_CASES_3                                # noqa: E402

DEFAULT_ADAA_RUN = os.path.join(_ROOT, "eval", "results", "full_2026_10_06_23_35_25")
LLM_RESULTS = os.path.join(_ROOT, "llm_codegen_experiment", "results")
DEFAULT_OUT = os.path.join(LLM_RESULTS, "comparison_adaa_vs_llm.xlsx")

# ---------------------------------------------------------------------------
# Query metadata: category, compare mode, tolerance, question text
# ---------------------------------------------------------------------------

_R1_GROUPS = ["Simple aggregation", "Filtering", "Grouping + ranking", "Time based", "Multi-condition",
              "Derived calculations"]
_COMPOUND_TAGS = {
    "two_scalars": "Compound: two scalars", "scalar_plus_ranked": "Compound: scalar + ranked list",
    "two_grouped_outputs": "Compound: two grouped outputs", "filter_then_two_agg": "Compound: filter + two aggregates",
}


def case_registry() -> Dict[str, Dict[str, Any]]:
    """id -> {query, category, compare_mode, float_tol, kind}."""
    reg: Dict[str, Dict[str, Any]] = {}
    for c in TEST_CASES + TEST_CASES_2 + TEST_CASES_3:
        n = int(c.id[1:])
        tags = set(c.tags)
        if n <= 30:
            category = f"{(n - 1) // 5 + 1}. {_R1_GROUPS[(n - 1) // 5]}"
        elif "pseudo_compound" in tags:
            category = "7. Pseudo-compound: comparison" if "comparison" in tags else "7. Pseudo-compound: breakdown"
        else:
            category = "8. " + next((v for k, v in _COMPOUND_TAGS.items() if k in tags), "Compound")
        reg[c.id] = {"query": c.query, "category": category, "compare_mode": c.compare_mode,
                     "float_tol": c.float_tol, "kind": "single"}
    for c in MULTITURN_TEST_CASES:
        last = c.turns[-1]
        reg[c.id] = {"query": last.query, "category": f"9. Multi-turn ({len(c.turns)} turns)",
                     "compare_mode": last.compare_mode, "float_tol": last.float_tol, "kind": "multi"}
    return reg


# ---------------------------------------------------------------------------
# Value comparison (column names and column order ignored)
# ---------------------------------------------------------------------------

def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _close(a: float, b: float, tol: float) -> bool:
    return abs(a - b) / max(abs(a), abs(b), 1e-10) <= tol


def _row_signature(row: Dict[str, Any]):
    nums = sorted(float(v) for v in row.values() if _is_num(v))
    strs = sorted(str(v).strip().lower() for v in row.values() if not _is_num(v))
    return nums, strs


def first_difference(gt: List[Dict[str, Any]], got: List[Dict[str, Any]], tol: float) -> Optional[str]:
    """None if the two tables hold the same values in the same row order; otherwise a short description."""
    if len(gt) != len(got):
        return f"{len(got)} rows returned, ground truth has {len(gt)}"
    for i, (g, p) in enumerate(zip(gt, got), start=1):
        gn, gs = _row_signature(g)
        pn, ps = _row_signature(p)
        if gs != ps or len(gn) != len(pn):
            return f"row {i}: {', '.join(map(str, gs + gn))} vs {', '.join(map(str, ps + pn))}"
        for a, b in zip(gn, pn):
            if not _close(a, b, tol):
                return f"row {i}: {a:,.2f} vs {b:,.2f}"
    return None


def _rows_equal(g: Dict[str, Any], p: Dict[str, Any], tol: float) -> bool:
    gn, gs = _row_signature(g)
    pn, ps = _row_signature(p)
    return gs == ps and len(gn) == len(pn) and all(_close(a, b, tol) for a, b in zip(gn, pn))


def rows_contained(gt: List[Dict[str, Any]], got: List[Dict[str, Any]], tol: float, ordered: bool) -> bool:
    """Every ground-truth row appears in the answer (extra rows allowed), names ignored. ordered=True keeps the row order."""
    if not gt or not got:
        return False
    if ordered:                       # the ground-truth rows must appear in the same relative order
        j = 0
        for p in got:
            if j < len(gt) and _rows_equal(gt[j], p, tol):
                j += 1
        return j == len(gt)
    used = set()                      # any order: match each ground-truth row to its own answer row
    for g in gt:
        hit = next((i for i, p in enumerate(got) if i not in used and _rows_equal(g, p, tol)), None)
        if hit is None:
            return False
        used.add(hit)
    return True


def name_insensitive_match(rec: Dict[str, Any], meta: Dict[str, Any]) -> bool:
    """
    Match against ground truth (column names ignored, extra rows allowed). A case the original comparator
    accepted stays correct; otherwise every ground-truth row must be present in the answer.
    """
    if rec.get("value_match"):
        return True
    return rows_contained(rec.get("gt_data") or [], rec.get("pipeline_data") or [], meta["float_tol"],
                          ordered=meta["compare_mode"] != "value_only")


def extra_rows(rec: Dict[str, Any]) -> int:
    gt, got = rec.get("gt_data") or [], rec.get("pipeline_data") or []
    return max(len(got) - len(gt), 0) if gt and got else 0


def render_value(rows: Optional[List[Dict[str, Any]]], max_rows: int = 5) -> str:
    """Short text for a table: a lone number as is, otherwise 'col=value | col=value' per row (first rows only)."""
    if not rows:
        return ""
    def fmt(v: Any) -> str:
        return f"{v:,.2f}" if isinstance(v, float) else str(v)
    if len(rows) == 1 and len(rows[0]) == 1:
        return fmt(next(iter(rows[0].values())))
    lines = [" | ".join(f"{k}={fmt(v)}" for k, v in r.items()) for r in rows[:max_rows]]
    if len(rows) > max_rows:
        lines.append(f"... (+{len(rows) - max_rows} more rows)")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Classification of one system's result for one query
# ---------------------------------------------------------------------------

def _sandbox_status(rec: Dict[str, Any]) -> str:
    return (rec.get("trace") or [{}])[0].get("sandbox_status", "") if rec.get("trace") else ""


def classify(rec: Dict[str, Any], meta: Dict[str, Any], system: str, adaa_label: Optional[str] = None) -> Dict[str, Any]:
    """
    outcome: correct | wrong_answer | error | refused_correctly | refused_wrongly | answered_should_refuse
    Returns {outcome, match, is_error, reason, value}.
    """
    expects_refuse = rec.get("expected_behavior") == "refuse"
    status = rec.get("pipeline_status")
    message = str(rec.get("pipeline_message") or "")
    skipped = rec.get("outcome") == "skipped"
    out = {"outcome": "", "match": False, "is_error": False, "reason": "", "value": ""}

    if skipped or status == "error" or (status is None and not rec.get("pipeline_data")):
        out.update(outcome="error", is_error=True, value="(error)")
        if skipped:
            out["reason"] = "context turn failed: " + "; ".join(str(m) for m in (rec.get("mismatches") or []))[:200]
        else:
            tag = _sandbox_status(rec) if system == "llm" else ""
            out["reason"] = (f"[{tag}] " if tag else "") + (message or "no usable answer")
        if adaa_label:
            out["reason"] = f"{adaa_label} | " + out["reason"]
        return out

    if status == "unsolvable":
        out["value"] = f"(refused: {message[:80]})"
        if expects_refuse:
            out.update(outcome="refused_correctly", match=True)
        else:
            out.update(outcome="refused_wrongly", match=False,
                       reason=(adaa_label + " | " if adaa_label else "") + f"refused an answerable query: {message[:140]}")
        return out

    out["value"] = render_value(rec.get("pipeline_data"))
    if expects_refuse:
        out.update(outcome="answered_should_refuse", match=False,
                   reason=(adaa_label + " | " if adaa_label else "")
                   + "answered a question that has no single-table answer (should have been refused)")
        return out

    if name_insensitive_match(rec, meta):
        out.update(outcome="correct", match=True)
    else:
        diff = first_difference(rec.get("gt_data") or [], rec.get("pipeline_data") or [], meta["float_tol"])
        out.update(outcome="wrong_answer", match=False,
                   reason=(adaa_label + " | " if adaa_label else "") + f"wrong value: {diff or 'values differ'}")
    return out


# ---------------------------------------------------------------------------
# Build the table
# ---------------------------------------------------------------------------

def _label_text(entry: Optional[Dict[str, Any]]) -> Optional[str]:
    if not entry:
        return None
    lbl = next((l for l in entry["labels"] if l["dimension"] == "outcome"), None)
    return lbl["label"] if lbl else None


def build_rows(adaa_recs: List[Dict[str, Any]], llm_recs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    reg = case_registry()
    adaa = {r["id"]: r for r in adaa_recs}
    llm = {r["id"]: r for r in llm_recs}
    labels = {e["id"]: e for e in fl.label_run(adaa_recs, fl.load_overrides())}
    rows = []
    for cid in sorted(reg, key=lambda i: (i.startswith("MT"), i)):
        meta = reg[cid]
        a, l = adaa[cid], llm[cid]
        ca = classify(a, meta, "adaa", _label_text(labels.get(cid)) if not a.get("passed") else None)
        cl = classify(l, meta, "llm")
        expects_refuse = a.get("expected_behavior") == "refuse"
        gt_text = "(correct = refuse: needs two separate results)" if expects_refuse else render_value(a.get("gt_data"))
        rows.append({
            "query_id": cid, "query": meta["query"], "category": meta["category"],
            "expected_behavior": a.get("expected_behavior"), "ground_truth_value": gt_text,
            "adaa_value": ca["value"], "llm_value": cl["value"],
            "gt_vs_adaa": ca["match"], "gt_vs_llm": cl["match"],
            "adaa_outcome": ca["outcome"], "llm_outcome": cl["outcome"],
            "adaa_error": ca["is_error"], "llm_error": cl["is_error"],
            "adaa_error_reason": ca["reason"], "llm_error_reason": cl["reason"],
            "adaa_extra_rows": extra_rows(a) if ca["match"] and not expects_refuse else 0,
            "llm_extra_rows": extra_rows(l) if cl["match"] and not expects_refuse else 0,
            "adaa_strict_match": bool(a.get("passed")), "llm_strict_match": bool(l.get("passed")),
            "adaa_cost_usd": a.get("cost_usd") or 0.0, "llm_cost_usd": l.get("cost_usd") or 0.0,
            "adaa_seconds": a.get("duration_s") or 0.0, "llm_seconds": l.get("duration_s") or 0.0,
        })
    return rows


# ---------------------------------------------------------------------------
# Summary tables
# ---------------------------------------------------------------------------

OUTCOMES = ["correct", "wrong_answer", "error", "refused_correctly", "refused_wrongly", "answered_should_refuse"]


def build_summary(rows: List[Dict[str, Any]], adaa_recs: List[Dict[str, Any]], llm_recs: List[Dict[str, Any]]) -> Dict[str, Any]:
    n = len(rows)
    answer_rows = [r for r in rows if r["expected_behavior"] == "answer"]
    refuse_rows = [r for r in rows if r["expected_behavior"] == "refuse"]

    def count(system: str, rs: List[Dict[str, Any]]) -> Dict[str, int]:
        c = Counter(r[f"{system}_outcome"] for r in rs)
        return {o: c.get(o, 0) for o in OUTCOMES}

    by_cat: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        d = by_cat.setdefault(r["category"], {"queries": 0, "adaa_correct": 0, "llm_correct": 0})
        d["queries"] += 1
        d["adaa_correct"] += int(r["gt_vs_adaa"])
        d["llm_correct"] += int(r["gt_vs_llm"])

    both = sum(1 for r in rows if r["gt_vs_adaa"] and r["gt_vs_llm"])
    only_adaa = sum(1 for r in rows if r["gt_vs_adaa"] and not r["gt_vs_llm"])
    only_llm = sum(1 for r in rows if r["gt_vs_llm"] and not r["gt_vs_adaa"])
    neither = n - both - only_adaa - only_llm

    def totals(recs: List[Dict[str, Any]]) -> Dict[str, float]:
        return {"cost_usd": sum(r.get("cost_usd") or 0 for r in recs), "seconds": sum(r.get("duration_s") or 0 for r in recs),
                "input_tokens": sum(r.get("input_tokens") or 0 for r in recs),
                "output_tokens": sum(r.get("output_tokens") or 0 for r in recs),
                "llm_calls": sum(r.get("llm_call_count") or 0 for r in recs), "queries": len(recs)}

    return {
        "n": n, "n_answer": len(answer_rows), "n_refuse": len(refuse_rows),
        "all": {"adaa": count("adaa", rows), "llm": count("llm", rows)},
        "answerable": {"adaa": count("adaa", answer_rows), "llm": count("llm", answer_rows)},
        "refuse": {"adaa": count("adaa", refuse_rows), "llm": count("llm", refuse_rows)},
        "match_all": {"adaa": sum(r["gt_vs_adaa"] for r in rows), "llm": sum(r["gt_vs_llm"] for r in rows)},
        "match_answerable": {"adaa": sum(r["gt_vs_adaa"] for r in answer_rows), "llm": sum(r["gt_vs_llm"] for r in answer_rows)},
        "correct_thanks_to_extra_rows": {"adaa": sum(1 for r in rows if r["adaa_extra_rows"] > 0),
                                         "llm": sum(1 for r in rows if r["llm_extra_rows"] > 0)},
        "strict_all": {"adaa": sum(r["adaa_strict_match"] for r in rows), "llm": sum(r["llm_strict_match"] for r in rows)},
        "by_category": by_cat,
        "agreement": {"both correct": both, "only ADAA correct": only_adaa, "only code system correct": only_llm,
                      "both not correct": neither},
        "adaa_totals": totals(adaa_recs), "llm_totals": totals(llm_recs),
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

COLUMNS = [
    ("query_id", 9), ("query", 52), ("category", 34), ("expected_behavior", 11), ("ground_truth_value", 38),
    ("adaa_value", 38), ("llm_value", 38), ("gt_vs_adaa", 11), ("gt_vs_llm", 11),
    ("adaa_outcome", 20), ("llm_outcome", 20), ("adaa_error", 10), ("llm_error", 10),
    ("adaa_error_reason", 60), ("llm_error_reason", 60),
    ("adaa_extra_rows", 11), ("llm_extra_rows", 11), ("adaa_strict_match", 12), ("llm_strict_match", 12),
    ("adaa_cost_usd", 11), ("llm_cost_usd", 11), ("adaa_seconds", 11), ("llm_seconds", 11),
]
_MATCH_COLS = {"gt_vs_adaa", "gt_vs_llm"}
_ERROR_COLS = {"adaa_error", "llm_error"}


def write_excel(path: str, rows: List[Dict[str, Any]], summary: Dict[str, Any], notes: List[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    bold_white = Font(bold=True, color="FFFFFF")

    # --- comparison sheet
    ws = wb.active
    ws.title = "comparison"
    for j, (name, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = bold_white, head_fill
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, r in enumerate(rows, start=2):
        for j, (name, _) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=i, column=j, value=r[name])
            c.alignment = Alignment(vertical="top", wrap_text=name in ("query", "ground_truth_value", "adaa_value", "llm_value",
                                                                       "adaa_error_reason", "llm_error_reason"))
            if name in _MATCH_COLS:
                c.fill = green if r[name] else red
            elif name in _ERROR_COLS and r[name]:
                c.fill = red
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # --- summary sheet
    sm = wb.create_sheet("summary")
    row = 1

    def put(values: List[Any], bold: bool = False, fill=None) -> None:
        nonlocal row
        for j, v in enumerate(values, start=1):
            c = sm.cell(row=row, column=j, value=v)
            if bold:
                c.font = bold_white if fill is head_fill else Font(bold=True)
            if fill is not None:
                c.fill = fill
        row += 1

    def table(title: str, header: List[str], body: List[List[Any]]) -> None:
        nonlocal row
        put([title], bold=True)
        put(header, bold=True, fill=head_fill)
        for b in body:
            put(b)
        row += 1

    s = summary
    table(f"Matches ground truth (column names ignored, extra rows allowed): of all {s['n']} queries and of the {s['n_answer']} answerable ones",
          ["", "ADAA", "Code-writing system"],
          [[f"All {s['n']} queries", s["match_all"]["adaa"], s["match_all"]["llm"]],
           [f"Answerable only ({s['n_answer']})", s["match_answerable"]["adaa"], s["match_answerable"]["llm"]],
           [f"Compound, correct = refuse ({s['n_refuse']}): refused correctly",
            s["refuse"]["adaa"]["refused_correctly"], s["refuse"]["llm"]["refused_correctly"]],
           [f"  of which correct only because extra rows are allowed", s["correct_thanks_to_extra_rows"]["adaa"],
            s["correct_thanks_to_extra_rows"]["llm"]],
           [f"All {s['n']} under STRICT scoring (column names and row count must match)", s["strict_all"]["adaa"], s["strict_all"]["llm"]]])
    table("Outcome counts, answerable queries only", ["outcome", "ADAA", "Code-writing system"],
          [[o, s["answerable"]["adaa"][o], s["answerable"]["llm"][o]] for o in OUTCOMES[:5]])
    table(f"Outcome counts, compound queries (correct = refuse), {s['n_refuse']} queries", ["outcome", "ADAA", "Code-writing system"],
          [[o, s["refuse"]["adaa"][o], s["refuse"]["llm"][o]] for o in ("refused_correctly", "answered_should_refuse", "error")])
    table("Who is right where (all queries)", ["", "queries"], [[k, v] for k, v in s["agreement"].items()])
    table("By category", ["category", "queries", "ADAA correct", "Code system correct"],
          [[k, v["queries"], v["adaa_correct"], v["llm_correct"]] for k, v in sorted(s["by_category"].items())])
    a, l = s["adaa_totals"], s["llm_totals"]
    table("Cost and time", ["", "ADAA", "Code-writing system"],
          [["total cost (USD)", round(a["cost_usd"], 5), round(l["cost_usd"], 5)],
           ["cost per query (USD)", round(a["cost_usd"] / max(a["queries"], 1), 5), round(l["cost_usd"] / max(l["queries"], 1), 5)],
           ["total time (s)", round(a["seconds"]), round(l["seconds"])],
           ["time per query (s)", round(a["seconds"] / max(a["queries"], 1), 1), round(l["seconds"] / max(l["queries"], 1), 1)],
           ["input tokens", a["input_tokens"], l["input_tokens"]], ["output tokens", a["output_tokens"], l["output_tokens"]],
           ["LLM calls", a["llm_calls"], l["llm_calls"]]])
    sm.column_dimensions["A"].width = 70
    sm.column_dimensions["B"].width = 18
    sm.column_dimensions["C"].width = 22
    sm.column_dimensions["D"].width = 22

    # --- notes sheet
    nt = wb.create_sheet("notes")
    nt.column_dimensions["A"].width = 160
    for i, line in enumerate(notes, start=1):
        c = nt.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if line and line == line.upper() and len(line) < 60:
            c.font = Font(bold=True)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def _manifest(run_dir: str) -> Dict[str, Any]:
    with open(os.path.join(run_dir, "manifest.json"), encoding="utf-8") as f:
        return json.load(f)


def build_notes(adaa_dir: str, llm_dir: str) -> List[str]:
    ma, ml = _manifest(adaa_dir), _manifest(llm_dir)
    return [
        "WHAT THIS IS",
        "One row per query (63): ADAA (LLM plans, fixed tools run, critic checks, LLM writes the sentence) versus a code-writing system "
        "(one LLM call writes pandas code, a restricted subprocess runs it). Both use gpt-oss-20b. Each system was run once on all 63 queries.",
        "",
        "SOURCE RUNS",
        f"ADAA: {os.path.relpath(adaa_dir, _ROOT)}  (commit {str(ma['git'].get('commit'))[:10]}, status {ma['status']}, {ma['totals']['records']} records)",
        f"Code-writing system: {os.path.relpath(llm_dir, _ROOT)}  (commit {str(ml['git'].get('commit'))[:10]}, prompt "
        f"{(ml.get('codegen') or {}).get('version')}, status {ml['status']}, {ml['totals']['records']} records)",
        f"Prices as of {ma['prices']['as_of']} (gpt-oss-20b: $0.075 in / $0.30 out per 1M tokens).",
        "",
        "COLUMN DEFINITIONS",
        "gt_vs_adaa / gt_vs_llm: TRUE when every ground-truth row is present in the system's table (1% relative tolerance). EXTRA ROWS ARE ALLOWED "
        "(for example 4 correct region rows plus a Total row is correct; decided by the user, applied to both systems). Column names and column order "
        "are ignored; row order must be kept for ranked questions. For the 10 compound queries (Q38-Q47, no ground truth) TRUE means the system refused.",
        "adaa_extra_rows / llm_extra_rows: how many rows beyond ground truth a correct answer contained (0 if none).",
        "adaa_error / llm_error: TRUE only when the system produced NO usable answer (crash, blocked by the sandbox, tool error, skipped).",
        "A WRONG ANSWER is not an error: it has error = FALSE, match = FALSE and outcome = wrong_answer. A wrong refusal is not an error either "
        "(outcome = refused_wrongly).",
        "outcome values: correct, wrong_answer, error, refused_correctly, refused_wrongly, answered_should_refuse.",
        "adaa_error_reason: ADAA's Part C failure category (from the failure-cause labeller) plus the detail. llm_error_reason: sandbox status "
        "and message, or the first differing value, or the refusal.",
        "adaa_strict_match / llm_strict_match: the original harness score, where column names and the row count must also match. Kept for transparency.",
        "",
        "CAVEATS",
        "One run per system: ADAA varies run to run (see the B3 consistency findings in docs/v2-polish-plan.md), so small differences are not significant.",
        "ADAA's cost and time include an answer-writing LLM call and a larger planner prompt; the code-writing system returns a table only, with no sentence.",
        "The code-writing system's sandbox is a restricted exec, adequate for this experiment on a public dataset but not safe for untrusted users.",
        "The code system was given the same one-table contract and an explicit REFUSE option as ADAA's planner has (prompt version v2).",
        "The sandbox allows only plainly safe df.query strings; anything else on its deny list is blocked and counts as an error.",
    ]


def _latest(pattern: str) -> Optional[str]:
    found = sorted(glob.glob(pattern))
    return found[-1] if found else None


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="ADAA vs code-writing system comparison table (Excel)")
    ap.add_argument("--adaa", default=DEFAULT_ADAA_RUN, help="ADAA full-eval run folder")
    ap.add_argument("--llm", default=_latest(os.path.join(LLM_RESULTS, "full_*")), help="code-writing run folder")
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args(argv)
    if not args.llm or not os.path.isdir(args.llm):
        print("No code-writing run folder found. Run llm_codegen_experiment/run_experiment.py first.")
        return 2
    adaa_recs = final_records(load_jsonl(os.path.join(args.adaa, "records.jsonl")))
    llm_recs = final_records(load_jsonl(os.path.join(args.llm, "records.jsonl")))
    rows = build_rows(adaa_recs, llm_recs)
    summary = build_summary(rows, adaa_recs, llm_recs)
    write_excel(args.out, rows, summary, build_notes(args.adaa, args.llm))
    s = summary
    print(f"Wrote {args.out}  ({len(rows)} rows)")
    print(f"Matches ground truth, all {s['n']}: ADAA {s['match_all']['adaa']}, code system {s['match_all']['llm']}  "
          f"| answerable only ({s['n_answer']}): ADAA {s['match_answerable']['adaa']}, code system {s['match_answerable']['llm']}")
    print("Agreement:", s["agreement"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
